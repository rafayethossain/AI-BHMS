"""
Tech-Pack Excel import (RQ-038).

Parses an uploaded — possibly hand-edited — tech-pack workbook back into a
:class:`TechPackDocument` DTO, tolerant of:

* reordered / renamed columns (keyword-based header detection, never position);
* a title row above the headers, blank spacer rows, and extra columns;
* empty cells (read back as ``""`` / ``None`` qty) and empty ``Match`` values;
* dates as real date cells *or* ``22/Mar/2022`` / ISO strings;
* duplicate style numbers (surfaced as a warning, never a failure).

A missing ``BOM`` sheet yields a structured ``errors`` entry (never a crash);
non-workbook input raises ``ValueError``.

Pure functions — no Django model dependency.
"""
from __future__ import annotations

import io
import os
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook

from .columns import BOM_KEYWORDS, DESIGN_INFO_KEYWORDS, clean_cell, find_header_row, map_columns
from .pdf_parser import TechPackBOMRow, TechPackDesignInfo, TechPackDocument

DESIGN_HEADER_MIN_HITS = 4  # enough for "Issue" + "Style" + a couple more
BOM_HEADER_MIN_HITS = 3


def parse_techpack_workbook(file: bytes | str | os.PathLike[str] | BinaryIO) -> TechPackDocument:
    """Parse a tech-pack workbook (bytes, path, or file-like) into a DTO.

    Raises ``ValueError`` for non-workbook input.
    """
    stream = _as_stream(file)
    try:
        workbook = load_workbook(stream, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl/BadZipFile on bad input
        raise ValueError(f"Not a valid Excel workbook: {exc}") from exc

    errors: list[str] = []
    warnings: list[str] = []

    design_sheet = _find_sheet(workbook, ("design", "infromation"))
    bom_sheet = _find_sheet(workbook, ("bom",))
    if bom_sheet is None:
        errors.append("Missing BOM sheet — expected a sheet titled like 'BOM'.")

    design_info = _read_design_info(design_sheet, warnings) if design_sheet is not None else TechPackDesignInfo()
    bom_rows = _read_bom(bom_sheet, warnings) if bom_sheet is not None else ()

    return TechPackDocument(
        design_info=design_info,
        bom_rows=bom_rows,
        errors=tuple(errors),
        warnings=tuple(warnings),
    )


# ---------------------------------------------------------------------- sheets


def _as_stream(file: bytes | str | os.PathLike[str] | BinaryIO) -> BinaryIO:
    if isinstance(file, bytes):
        return io.BytesIO(file)
    if isinstance(file, (str, os.PathLike)):
        return open(Path(file), "rb")  # noqa: SIM115 — owned by load_workbook
    return file  # file-like; caller owns closing


def _find_sheet(workbook, keywords: tuple[str, ...]):
    for sheet in workbook.worksheets:
        title = sheet.title.lower()
        if any(keyword in title for keyword in keywords):
            return sheet
    return None


# ---------------------------------------------------------------- design info


def _read_design_info(ws, warnings: list[str]) -> TechPackDesignInfo:
    header_row = find_header_row(ws, DESIGN_INFO_KEYWORDS, DESIGN_HEADER_MIN_HITS)
    if header_row is None:
        warnings.append("Design Infromation header row not found.")
        return TechPackDesignInfo()

    columns = map_columns(
        {cell.column: cell.value for cell in ws[header_row]},
        DESIGN_INFO_KEYWORDS,
    )
    value_row = header_row + 1

    def cell(field: str) -> Any:
        col = columns[field]
        return ws.cell(row=value_row, column=col).value if col is not None else None

    issue_date = _parse_date(cell("issue_date"), warnings)
    fields = {field: clean_cell(cell(field)) for field in DESIGN_INFO_KEYWORDS}
    fields["issue_date"] = issue_date
    fields["note"] = _clean_note(cell("note"))

    _flag_duplicate_style_numbers(ws, header_row, columns["style_number"], warnings)

    return TechPackDesignInfo(**fields)


def _flag_duplicate_style_numbers(ws, header_row: int, style_col: int | None, warnings: list[str]) -> None:
    if style_col is None:
        return
    seen: set[str] = set()
    for row in range(header_row + 1, ws.max_row + 1):
        value = clean_cell(ws.cell(row=row, column=style_col).value)
        if value:
            seen.add(value)
    if len(seen) > 1:
        ordered = ", ".join(sorted(seen))
        warnings.append(f"Multiple style numbers found in workbook ({ordered}) — confirm the intended one.")


# -------------------------------------------------------------------------- BOM


def _read_bom(ws, warnings: list[str]) -> tuple[TechPackBOMRow, ...]:
    header_row = find_header_row(ws, BOM_KEYWORDS, BOM_HEADER_MIN_HITS)
    if header_row is None:
        warnings.append("BOM header row not found in the BOM sheet.")
        return ()

    columns = map_columns({cell.column: cell.value for cell in ws[header_row]}, BOM_KEYWORDS)
    rows: list[TechPackBOMRow] = []
    for row_index in range(header_row + 1, ws.max_row + 1):
        values = {
            field: clean_cell(ws.cell(row=row_index, column=col).value)
            if col is not None else ""
            for field, col in columns.items()
        }
        if not values["type"]:
            continue  # spacer / stray rows
        if not values["description_code"] and not values["location"] and not values["supplier"]:
            continue  # a row that carries only a type is not a real BOM line
        rows.append(
            TechPackBOMRow(
                type=values["type"],
                description_code=values["description_code"],
                location=values["location"],
                supplier=values["supplier"],
                colour=values["colour"],
                width_size=values["width_size"],
                qty=_parse_qty(ws.cell(row=row_index, column=columns["qty"]).value, warnings)
                if columns["qty"] is not None else None,
                match=values["match"],
            )
        )
    return tuple(rows)


# ------------------------------------------------------------------- coercion


def _clean_note(value: Any) -> str:
    """Normalise the note cell, preserving its line breaks (the PDF note is
    multi-line; ``clean_cell`` would collapse it into one space-joined run)."""
    if value is None:
        return ""
    lines = [re.sub(r"[ \t]+", " ", str(line)).strip() for line in str(value).split("\n")]
    return "\n".join(line for line in lines if line)


def _parse_date(value: Any, warnings: list[str]) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value or not isinstance(value, str):
        return None
    raw = value.strip()
    for fmt in ("%d/%b/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    warnings.append(f"Unrecognised issue-date format: {value!r}")
    return None


def _parse_qty(value: Any, warnings: list[str]) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value).strip())
    except InvalidOperation:
        warnings.append(f"BOM qty could not be parsed: {value!r}")
        return None
