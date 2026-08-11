"""
RQ-038 — Tech-Pack Excel Import Service tests.

Covers `apps.merchandising.techpack.excel_import.parse_techpack_workbook`:
keyword-based header detection (tolerant of reordered/renamed columns and
extra rows), date/Decimal coercion, structured errors for missing BOM sheet,
empty-cell handling, duplicate style-number surfacing, and lossless round-trip
through the RQ-037 writer and the golden sample workbook.

Pure functions — no Django DB.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from apps.merchandising.techpack.excel_export import write_techpack_workbook
from apps.merchandising.techpack.excel_import import parse_techpack_workbook
from apps.merchandising.techpack.pdf_parser import (
    StyleTechPackParser,
    TechPackDocument,
)

PDF_EXTRACT_DIR = Path(__file__).resolve().parents[2].parent / "PDF Extract"
SAMPLE_PDF = PDF_EXTRACT_DIR / "Sample style doc.pdf"
SAMPLE_XLSX = PDF_EXTRACT_DIR / "Extracted_2026-07-13 .xlsx"

DESIGN_HEADERS = [
    "Issue Date", "Block", "Based On", "Customer", "Style Number", "Size",
    "Designer", "Pattern Cutter", "Issuer", "Cloth Code", "Length", "Sketch",
    "Description", "Note",
]

BOM_HEADERS = [
    "Type", "Description/ Code ", "Location", "Supplier", "Colour",
    "Width/Size", "Qty", "Match",
]


def _build_workbook(design_headers, design_values, bom_headers, bom_rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Design Infromation"
    ws.append(design_headers)
    ws.append(design_values)
    ws2 = wb.create_sheet("BOM")
    ws2.append(bom_headers)
    for row in bom_rows:
        ws2.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def parsed_doc() -> TechPackDocument:
    return StyleTechPackParser().parse(SAMPLE_PDF.read_bytes())


@pytest.fixture(scope="module")
def writer_roundtrip(parsed_doc) -> TechPackDocument:
    """RQ-037 output fed back through RQ-038."""
    data = write_techpack_workbook(parsed_doc).read()
    return parse_techpack_workbook(data)


# --------------------------------------------------------- round-trip (writer)


def test_writer_output_round_trips_to_dto(writer_roundtrip, parsed_doc):
    assert writer_roundtrip.to_dict() == parsed_doc.to_dict()


def test_issue_date_survives_round_trip(writer_roundtrip):
    assert writer_roundtrip.design_info.issue_date == date(2022, 3, 22)


def test_bom_qty_survives_as_decimal(writer_roundtrip):
    assert writer_roundtrip.bom_rows[0].qty == Decimal("1.67")
    assert all(isinstance(row.qty, Decimal) for row in writer_roundtrip.bom_rows)


def test_empty_match_column_reads_as_empty_string(writer_roundtrip):
    assert all(row.match == "" for row in writer_roundtrip.bom_rows)


# ---------------------------------------------------- golden sample workbook


def test_sample_workbook_parses_without_error():
    doc = parse_techpack_workbook(SAMPLE_XLSX.read_bytes())
    assert doc.errors == ()
    assert len(doc.bom_rows) == 6
    assert doc.design_info.issue_date == date(2026, 3, 22)  # sample's hand-edit
    assert doc.design_info.style_number == "67741T"
    assert doc.design_info.description == "565235 LB LIZZIE WIDE LEG PANT"


# -------------------------------------------------- keyword header detection


def test_headers_detected_by_keyword_not_position():
    # Reordered columns + a title row above + a trailing spacer row.
    wb = Workbook()
    ws = wb.active
    ws.title = "Design Infromation"
    ws.append(["CONFIDENTIAL"])
    ws.append(["Note", "Description", "Length", "Cloth Code", "Issuer",
               "Pattern Cutter", "Designer", "Size", "Style Number", "Customer",
               "Based On", "Block", "Issue Date", "Sketch"])
    ws.append(["hello note", "DESC", "0", "CLOTH", "Clone", "HAI", "Emmi",
               "10", "67741T", "DOTTI", "59073T", "59073T", "22/Mar/2022", ""])
    ws.append(["", "", "", "", "", "", "", "", "", "", "", "", "", ""])
    ws2 = wb.create_sheet("BOM")
    ws2.append(["Location", "Supplier", "Description/ Code ", "Type", "Qty",
                "Colour", "Width/Size", "Match"])
    ws2.append(["MAIN", "ALICE-", "SANDWASH LINEN XK-529", "Cloth", "1.67",
                "BLACK", "132 CM", ""])
    buf = BytesIO()
    wb.save(buf)
    doc = parse_techpack_workbook(buf.getvalue())
    assert doc.errors == ()
    assert doc.design_info.note == "hello note"
    assert doc.design_info.description == "DESC"
    assert doc.design_info.issue_date == date(2022, 3, 22)
    assert doc.design_info.style_number == "67741T"
    assert doc.design_info.cloth_code == "CLOTH"
    assert doc.bom_rows[0].type == "Cloth"
    assert doc.bom_rows[0].qty == Decimal("1.67")


def test_renamed_header_matches_keyword():
    wb = Workbook()
    ws = wb.active
    ws.title = "Design Infromation"
    ws.append(["Issue Date", "Block", "Based On", "Customer", "Style No.",
               "Size", "Designer", "Pattern Cutter", "Issuer", "Cloth",
               "Length", "Sketch", "Desc", "Note"])
    ws.append([datetime(2022, 3, 22), "59073T", "59073T", "DOTTI", "67741T",
               "10", "Emmi.Huynh", "HAI", "Clone", "SANDWASH LINEN", "0",
               "", "DESC", "note"])
    ws2 = wb.create_sheet("BOM")
    ws2.append(BOM_HEADERS)
    buf = BytesIO()
    wb.save(buf)
    doc = parse_techpack_workbook(buf.getvalue())
    assert doc.design_info.style_number == "67741T"
    assert doc.design_info.cloth_code == "SANDWASH LINEN"
    assert doc.design_info.description == "DESC"


# ------------------------------------------------------------- error handling


def test_missing_bom_sheet_returns_structured_error():
    wb = Workbook()
    ws = wb.active
    ws.title = "Design Infromation"
    ws.append(DESIGN_HEADERS)
    ws.append(["22/Mar/2022", "59073T", "59073T", "DOTTI", "67741T", "10",
               "Emmi.Huynh", "HAI", "Clone", "SANDWASH LINEN", "0", "",
               "DESC", ""])
    buf = BytesIO()
    wb.save(buf)
    doc = parse_techpack_workbook(buf.getvalue())
    assert any("BOM" in e for e in doc.errors)
    assert doc.bom_rows == ()


def test_invalid_workbook_raises_valueerror():
    with pytest.raises(ValueError):
        parse_techpack_workbook(b"this is not an xlsx")


# ------------------------------------------------------------- date / decimal


def test_string_date_coerced_to_date():
    doc = parse_techpack_workbook(
        _build_workbook(
            DESIGN_HEADERS,
            ["22/Mar/2022", "59073T", "59073T", "DOTTI", "67741T", "10",
             "Emmi.Huynh", "HAI", "Clone", "SANDWASH LINEN", "0", "", "D", ""],
            BOM_HEADERS,
            [],
        )
    )
    assert doc.design_info.issue_date == date(2022, 3, 22)


def test_iso_string_date_coerced_to_date():
    doc = parse_techpack_workbook(
        _build_workbook(
            DESIGN_HEADERS,
            ["2022-03-22", "59073T", "59073T", "DOTTI", "67741T", "10",
             "Emmi.Huynh", "HAI", "Clone", "SANDWASH LINEN", "0", "", "D", ""],
            BOM_HEADERS,
            [],
        )
    )
    assert doc.design_info.issue_date == date(2022, 3, 22)


def test_qty_string_coerced_to_decimal():
    doc = parse_techpack_workbook(
        _build_workbook(DESIGN_HEADERS, [None] * 14, BOM_HEADERS,
                        [["Cloth", "XK-529", "MAIN", "ALICE-", "BLACK",
                          "132 CM", "1.67", ""]]),
    )
    assert doc.bom_rows[0].qty == Decimal("1.67")


def test_qty_blank_becomes_none():
    doc = parse_techpack_workbook(
        _build_workbook(DESIGN_HEADERS, [None] * 14, BOM_HEADERS,
                        [["Cloth", "XK-529", "MAIN", "ALICE-", "BLACK",
                          "132 CM", "", ""]]),
    )
    assert doc.bom_rows[0].qty is None


def test_unparseable_date_warns_not_fails():
    doc = parse_techpack_workbook(
        _build_workbook(
            DESIGN_HEADERS,
            ["not-a-date", "59073T", "59073T", "DOTTI", "67741T", "10",
             "Emmi.Huynh", "HAI", "Clone", "SANDWASH LINEN", "0", "", "D", ""],
            BOM_HEADERS,
            [],
        )
    )
    assert doc.design_info.issue_date is None
    assert any("date" in w.lower() for w in doc.warnings)


# ------------------------------------------------------- empty / extra content


def test_empty_cells_coerce_to_empty_string(writer_roundtrip):
    dto = writer_roundtrip.design_info
    assert dto.sketch == ""
    assert dto.based_on == "59073T"


def test_extra_bom_rows_without_type_skipped():
    doc = parse_techpack_workbook(
        _build_workbook(DESIGN_HEADERS, [None] * 14, BOM_HEADERS,
                        [["Cloth", "XK-529", "MAIN", "ALICE-", "BLACK",
                          "132 CM", "1.67", ""],
                         [None, None, None, None, None, None, None, None],
                         ["", "", "STRAY", "", "", "", "", ""]]),
    )
    assert len(doc.bom_rows) == 1
    assert doc.bom_rows[0].type == "Cloth"


def test_extra_columns_ignored():
    doc = parse_techpack_workbook(
        _build_workbook(
            DESIGN_HEADERS + ["Extra"],
            ["22/Mar/2022", "59073T", "59073T", "DOTTI", "67741T", "10",
             "Emmi.Huynh", "HAI", "Clone", "SANDWASH LINEN", "0", "", "D", "",
             "junk"],
            BOM_HEADERS + ["Extra"],
            [["Cloth", "XK-529", "MAIN", "ALICE-", "BLACK", "132 CM",
              "1.67", "", "junk"]],
        )
    )
    assert doc.design_info.style_number == "67741T"
    assert len(doc.bom_rows) == 1


def test_duplicate_style_numbers_flagged_not_failed():
    wb = Workbook()
    ws = wb.active
    ws.title = "Design Infromation"
    ws.append(DESIGN_HEADERS)
    ws.append(["22/Mar/2022", "59073T", "59073T", "DOTTI", "67741T", "10",
               "Emmi.Huynh", "HAI", "Clone", "SANDWASH LINEN", "0", "", "D", ""])
    ws.append(["22/Mar/2022", "59073T", "59073T", "DOTTI", "99999X", "10",
               "Emmi.Huynh", "HAI", "Clone", "SANDWASH LINEN", "0", "", "D", ""])
    ws2 = wb.create_sheet("BOM")
    ws2.append(BOM_HEADERS)
    buf = BytesIO()
    wb.save(buf)
    doc = parse_techpack_workbook(buf.getvalue())
    assert any("style" in w.lower() and "67741T" in w for w in doc.warnings)
