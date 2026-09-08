"""
Design register Excel export API.

``GET /api/v1/merchandising/design-sheets/export/`` streams a tenant-scoped
.xlsx workbook (one row per design sheet) mirroring the register grid columns,
including the **Buyer** column, with display labels for status and
relationship. Requires ``merchandising:view`` (same as list).
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from rest_framework import status
from rest_framework.test import APIClient

from apps.merchandising.models import DesignSheet, StyleTechPack
from apps.setup.models import Buyer, ProductCategory, ProductType
from apps.tenants.models import Tenant
from apps.users.models import Permission, Role, RolePermission, UserRole

from openpyxl import load_workbook

User = get_user_model()


@pytest.fixture
def api_client():
    return APIClient()


@pytest.fixture
def tenant(db):
    return Tenant.objects.create(
        name="Export Co", slug="export-co",
        schema_name="tenant_exportco", status="active",
    )


@pytest.fixture
def other_tenant(db):
    return Tenant.objects.create(
        name="Other Co", slug="other-export-co",
        schema_name="tenant_export_other", status="active",
    )


def _role_with(t, name, permissions):
    role = Role.objects.create(tenant=t, name=name, is_system=True)
    for module, action in permissions:
        perm, _ = Permission.objects.get_or_create(
            module=module, action=action,
            defaults={"description": f"{module}:{action}"},
        )
        RolePermission.objects.create(role=role, permission=perm)
    return role


def _user(t, role, username, email):
    user = User.objects.create_user(
        username=username, email=email, password="testpass123!@#",
        tenant=t, status="active",
    )
    UserRole.objects.create(user=user, role=role)
    return user


@pytest.fixture
def user(tenant):
    role = _role_with(tenant, "ExportView", [("merchandising", "view")])
    return _user(tenant, role, "exportview", "exportview@test.com")


@pytest.fixture
def client(api_client, user):
    login = api_client.post("/api/v1/auth/login/", {
        "email": "exportview@test.com", "password": "testpass123!@#",
    })
    api_client.credentials(HTTP_AUTHORIZATION=f"Bearer {login.data['access']}")
    return api_client


@pytest.fixture
def no_perm_client(api_client, tenant):
    role = _role_with(tenant, "NoMech", [("setup", "view")])
    noperm = _user(tenant, role, "exportnoperm", "exportnoperm@test.com")
    login = api_client.post("/api/v1/auth/login/", {
        "email": "exportnoperm@test.com", "password": "testpass123!@#",
    })
    api_client.credentials(HTTP_AUTHORIZATION=f"Bearer {login.data['access']}")
    return api_client


@pytest.fixture
def category(tenant):
    return ProductCategory.objects.create(tenant=tenant, name="Apparel", code="APP")


@pytest.fixture
def product_type(tenant, category):
    return ProductType.objects.create(
        tenant=tenant, code="JGR", name="Jogger", category=category,
    )


@pytest.fixture
def buyer(tenant):
    return Buyer.objects.create(tenant=tenant, name="Prime Buyer", code="PB01")


@pytest.fixture
def sheet(tenant, buyer, product_type):
    tp = StyleTechPack.objects.create(
        tenant=tenant,
        techpack_number="TP-871",
        style_code="DS-1010",
        product_type=product_type,
        buyer=buyer,
        style_number="90123",
        block="59080T",
        based_on="59070T",
        description="Export me",
        designer="Emmi.Huynh",
        issuer="Issuer One",
        cloth_code="CC-100",
        contains="Div 3 / 3446",
        risk_date=date(2026, 9, 1),
        pattern_request_date=date(2026, 8, 15),
        relationship=StyleTechPack.Relationship.BASED_ON,
        customer="Prime Buyer",
    )
    return DesignSheet.objects.create(
        tenant=tenant,
        tech_pack=tp,
        status=DesignSheet.Status.NEW,
    )


@pytest.fixture
def foreign_sheet(other_tenant):
    tp = StyleTechPack.objects.create(
        tenant=other_tenant,
        techpack_number="TP-9999",
        style_code="DS-FOREIGN",
        customer="Other Buyer",
    )
    return DesignSheet.objects.create(
        tenant=other_tenant, tech_pack=tp, status=DesignSheet.Status.NEW,
    )


@pytest.mark.django_db
class TestDesignRegisterExport:
    def test_export_returns_tenant_scoped_xlsx_with_buyer_column(
        self, client, sheet, foreign_sheet
    ):
        resp = client.get("/api/v1/merchandising/design-sheets/export/")
        assert resp.status_code == status.HTTP_200_OK
        assert "spreadsheetml" in resp["Content-Type"]
        assert resp["Content-Disposition"].startswith("attachment")
        assert "design_register.xlsx" in resp["Content-Disposition"]

        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        assert ws.title == "Design Register"
        headers = [c.value for c in ws[1]]
        assert "File Number" in headers
        assert "Style Code" in headers
        assert "Buyer" in headers
        assert "Relationship" in headers
        assert "Status" in headers
        assert "Product Type" in headers
        assert "Product Category" in headers
        assert "Style Type" not in headers
        assert "Contains" not in headers

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert any(r[0] == "TP-871" for r in rows)
        data_row = next(r for r in rows if r[0] == "TP-871")
        idx = {h: i for i, h in enumerate(headers)}
        assert data_row[idx["Style Code"]] == "DS-1010"
        assert data_row[idx["Buyer"]] == "Prime Buyer"
        assert data_row[idx["Relationship"]] == "Based on"
        assert data_row[idx["Status"]] == "New"
        assert data_row[idx["Product Type"]] == "Jogger"
        assert data_row[idx["Product Category"]] == "Apparel"

        # Foreign tenant sheet is excluded (RBAC + tenant isolation)
        codes = [r[idx["Style Code"]] for r in rows]
        assert "DS-FOREIGN" not in codes

    def test_export_requires_merchandising_view_permission(self, no_perm_client):
        resp = no_perm_client.get("/api/v1/merchandising/design-sheets/export/")
        assert resp.status_code == status.HTTP_403_FORBIDDEN